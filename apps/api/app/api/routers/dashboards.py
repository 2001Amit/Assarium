from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from fastapi import APIRouter, Depends, Response
from pydantic import BaseModel, Field

from app.auth.deps import get_scope, get_tenant_context
from app.core.errors import NotFoundError, ValidationError
from app.dashboards.generator import generate_dashboard
from app.dashboards.runner import run_dashboard
from app.dashboards.types import Dashboard, DashboardData
from app.engine.factory import get_engine
from app.models.base import new_id
from app.models.entities import Connection, DashboardRecord, SemanticModelRecord
from app.semantic.types import Filter, SemanticModel
from app.tenancy.scope import TenantScope

# Every route below resolves a tenant before the handler runs. Declared once here
# rather than per endpoint, so a new endpoint is guarded by default instead of by
# whoever remembers. `assert_routes_are_guarded` at startup proves none slipped past.
router = APIRouter(
    prefix="/api", tags=["dashboards"],
    dependencies=[Depends(get_tenant_context)],
)


class DataRequest(BaseModel):
    """Filters chosen in the UI, applied on top of whatever the dashboard defines."""

    filters: list[Filter] = Field(default_factory=list)


class DashboardSummary(BaseModel):
    id: str
    connection_id: str
    name: str
    tile_count: int
    generated: bool


def _model(scope: TenantScope, connection_id: str) -> SemanticModel:
    TenantScope.assert_is_scope(scope, "_model")
    record = scope.db.execute(
        scope.select(SemanticModelRecord).where(SemanticModelRecord.connection_id == connection_id)
    ).scalar_one_or_none()
    if record is None:
        raise NotFoundError(
            "Build the semantic model first. Dashboards are made of its measures."
        )
    return SemanticModel.model_validate(record.document)


def _dashboard(record: DashboardRecord) -> Dashboard:
    return Dashboard.model_validate(record.document)


# --------------------------------------------------------------------------------------
# lifecycle
# --------------------------------------------------------------------------------------


@router.post("/connections/{connection_id}/dashboards/generate", response_model=Dashboard)
def generate(connection_id: str, scope: TenantScope = Depends(get_scope)) -> Dashboard:
    """Compose a dashboard from what the semantic model can actually support."""
    scope.get(Connection, connection_id)
    model = _model(scope, connection_id)

    dashboard = generate_dashboard(new_id(), model)
    scope.create(
        DashboardRecord,
        id=dashboard.id,
        connection_id=connection_id,
        name=dashboard.name,
        document=dashboard.model_dump(mode="json"),
        generated=True,
    )
    scope.db.commit()
    return dashboard


@router.get("/connections/{connection_id}/dashboards", response_model=list[DashboardSummary])
def list_dashboards(
    connection_id: str, scope: TenantScope = Depends(get_scope)
) -> list[DashboardSummary]:
    records = scope.db.execute(
        scope.select(DashboardRecord)
        .where(DashboardRecord.connection_id == connection_id)
        .order_by(DashboardRecord.created_at.desc())
    ).scalars()
    return [
        DashboardSummary(
            id=record.id,
            connection_id=record.connection_id,
            name=record.name,
            tile_count=len(record.document.get("tiles", [])),
            generated=record.generated,
        )
        for record in records
    ]


@router.get("/dashboards/{dashboard_id}", response_model=Dashboard)
def get_dashboard(dashboard_id: str, scope: TenantScope = Depends(get_scope)) -> Dashboard:
    return _dashboard(scope.get(DashboardRecord, dashboard_id))


@router.put("/dashboards/{dashboard_id}", response_model=Dashboard)
def save_dashboard(
    dashboard_id: str, dashboard: Dashboard, scope: TenantScope = Depends(get_scope)
) -> Dashboard:
    record = scope.get(DashboardRecord, dashboard_id)

    # Checked here rather than at render time, so a mis-specified tile is a refusal while
    # somebody is still editing it - not an empty card discovered at breakfast. Every
    # broken tile is reported at once; fixing them one round-trip at a time is worse.
    problems: list[str] = []
    for tile in dashboard.tiles:
        problems += [f"{tile.title or tile.id}: {reason}" for reason in tile.validate_shape()]
    if problems:
        raise ValidationError(
            "This dashboard has tiles that cannot be drawn as specified.",
            details={"tiles": problems},
        )

    dashboard.id = dashboard_id
    dashboard.connection_id = record.connection_id
    record.name = dashboard.name
    record.document = dashboard.model_dump(mode="json")
    record.generated = False
    scope.db.commit()
    return dashboard


@router.delete("/dashboards/{dashboard_id}", status_code=204)
def delete_dashboard(dashboard_id: str, scope: TenantScope = Depends(get_scope)) -> None:
    scope.db.delete(scope.get(DashboardRecord, dashboard_id))
    scope.db.commit()


# --------------------------------------------------------------------------------------
# data
# --------------------------------------------------------------------------------------


@router.post("/dashboards/{dashboard_id}/data", response_model=DashboardData)
def dashboard_data(
    dashboard_id: str, request: DataRequest, scope: TenantScope = Depends(get_scope)
) -> DashboardData:
    record = scope.get(DashboardRecord, dashboard_id)
    dashboard = _dashboard(record)
    model = _model(scope, record.connection_id)
    return run_dashboard(
        model, get_engine(scope.tenant_id), dashboard, request.filters,
        permissions=scope.context.permissions,
    )


# --------------------------------------------------------------------------------------
# export
# --------------------------------------------------------------------------------------


@router.post("/dashboards/{dashboard_id}/export.xlsx")
def export_xlsx(
    dashboard_id: str, request: DataRequest, scope: TenantScope = Depends(get_scope)
) -> Response:
    """
    Export the dashboard as a workbook.

    The first sheet carries the measure definitions and the filters that were applied,
    so the numbers travel with the reason they are what they are. A spreadsheet of
    figures with no definitions is how two teams end up arguing about revenue.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise NotFoundError(
            "Excel export needs the 'openpyxl' package: "
            "./.venv/bin/pip install openpyxl"
        ) from exc

    record = scope.get(DashboardRecord, dashboard_id)
    dashboard = _dashboard(record)
    model = _model(scope, record.connection_id)
    data = run_dashboard(
        model, get_engine(scope.tenant_id), dashboard, request.filters,
        permissions=scope.context.permissions,
    )

    workbook = Workbook()
    header_font = Font(bold=True, size=10)
    header_fill = PatternFill("solid", fgColor="EEEEF0")

    cover = workbook.active
    cover.title = "About"
    cover.column_dimensions["A"].width = 34
    cover.column_dimensions["B"].width = 68

    rows: list[tuple[str, str]] = [
        ("Dashboard", dashboard.name),
        ("Exported", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")),
        ("Engine", get_engine(scope.tenant_id).name),
        ("", ""),
        ("Filters applied", ""),
    ]
    applied = [*request.filters, *dashboard.filters]
    if applied:
        for filter_ in applied:
            values = ", ".join(str(v) for v in filter_.values)
            rows.append((f"   {filter_.field}", f"{filter_.operator} {values}".strip()))
    else:
        rows.append(("   None", "The whole dataset is included."))

    rows.append(("", ""))
    rows.append(("Measure definitions", ""))
    used: set[str] = set()
    for tile in dashboard.tiles:
        used.update(tile.measures)
    for measure_id in sorted(used):
        measure = model.measure(measure_id)
        if measure is None:
            continue
        source = (
            f"{measure.aggregation}({measure.column})"
            if measure.column
            else measure.aggregation
        )
        entity = model.entity(measure.entity_id)
        definition = f"{source} over {entity.label if entity else measure.entity_id}"
        if measure.description:
            definition += f" — {measure.description}"
        rows.append((f"   {measure.label}", definition))

    for index, (left, right) in enumerate(rows, start=1):
        cover.cell(row=index, column=1, value=left).font = Font(
            bold=left != "" and not left.startswith("   "), size=10
        )
        cover.cell(row=index, column=2, value=right).alignment = Alignment(wrap_text=True)

    used_titles: set[str] = set()
    for tile_result in data.tiles:
        if tile_result.error or (not tile_result.rows and not tile_result.stats):
            continue
        title = _sheet_name(tile_result.title, used_titles)
        sheet = workbook.create_sheet(title)

        if tile_result.type == "stat":
            sheet.append(["Measure", "Value", "Previous", "Change"])
            for stat in tile_result.stats:
                measure = model.measure(stat.measure)
                sheet.append(
                    [
                        measure.label if measure else stat.measure,
                        stat.value,
                        stat.previous,
                        stat.delta,
                    ]
                )
        else:
            sheet.append(tile_result.columns)
            for row in tile_result.rows:
                sheet.append([_cell(value) for value in row])

        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        sheet.freeze_panes = "A2"
        for column_index in range(1, sheet.max_column + 1):
            letter = get_column_letter(column_index)
            longest = max(
                (len(str(cell.value)) for cell in sheet[letter] if cell.value),
                default=10,
            )
            sheet.column_dimensions[letter].width = min(42, longest + 3)

    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = _safe_filename(dashboard.name) + ".xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/dashboards/{dashboard_id}/tiles/{tile_id}/export.csv")
def export_tile_csv(
    dashboard_id: str, tile_id: str, request: DataRequest, scope: TenantScope = Depends(get_scope)
) -> Response:
    import csv

    record = scope.get(DashboardRecord, dashboard_id)
    dashboard = _dashboard(record)
    model = _model(scope, record.connection_id)
    data = run_dashboard(
        model, get_engine(scope.tenant_id), dashboard, request.filters,
        permissions=scope.context.permissions,
    )

    tile_result = next((t for t in data.tiles if t.tile_id == tile_id), None)
    if tile_result is None:
        raise NotFoundError("That tile is not on this dashboard.")
    if tile_result.error:
        raise NotFoundError(tile_result.error)

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    if tile_result.type == "stat":
        writer.writerow(["measure", "value", "previous", "change"])
        for stat in tile_result.stats:
            writer.writerow([stat.measure, stat.value, stat.previous, stat.delta])
    else:
        writer.writerow(tile_result.columns)
        writer.writerows([[_cell(value) for value in row] for row in tile_result.rows])

    filename = _safe_filename(tile_result.title) + ".csv"
    return Response(
        content=buffer.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _cell(value: Any) -> Any:
    from decimal import Decimal

    if isinstance(value, Decimal):
        return float(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _sheet_name(title: str, used: set[str]) -> str:
    """Excel sheet names cap at 31 characters and must be unique."""
    cleaned = "".join(c for c in title if c not in "[]:*?/\\")[:31] or "Sheet"
    candidate = cleaned
    suffix = 2
    while candidate in used:
        candidate = f"{cleaned[:28]}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _safe_filename(value: str) -> str:
    cleaned = "".join(c if c.isalnum() or c in " -_" else "" for c in value).strip()
    return (cleaned or "assarium-export").replace(" ", "-").lower()[:60]
